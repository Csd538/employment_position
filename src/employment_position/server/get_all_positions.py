import asyncio
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from util.config import (
    POSITION_BATCH_INTERVAL_SECONDS,
    POSITION_BATCH_SIZE,
    POSITION_DETAIL_BATCH_SIZE,
    POSITION_DETAIL_CONCURRENCY,
    POSITION_DETAIL_REQUEST_INTERVAL_SECONDS,
    POSITION_MAX_PAGES,
    POSITION_OUTPUT_DIR,
    POSITION_PAGE_SIZE,
    PROVINCES,
)
from util.http_util import async_request_with_proxy


URL = (
    "https://ggfw.hrss.gd.gov.cn/recruitment/internet/main/"
    "internet/retrieval/c/recruitment/homepage/positions"
)
POSITION_DETAIL_URL = (
    "https://ggfw.hrss.gd.gov.cn/recruitment/internet/main/"
    "internet/r/c/webpage/homepage/position/detail/{position_id}"
)

HEADERS = {
    "Accept": "application/json, text/plain, */*",
    "Accept-Language": "zh-CN,zh;q=0.9",
    "Connection": "keep-alive",
    "Content-Type": "application/json",
    "Origin": "https://ggfw.hrss.gd.gov.cn",
    "Referer": "https://ggfw.hrss.gd.gov.cn/recruitment/internet/main/",
    "Sec-Fetch-Dest": "empty",
    "Sec-Fetch-Mode": "cors",
    "Sec-Fetch-Site": "same-origin",
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/153.0.0.0 Safari/537.36"
    ),
    "sec-ch-ua": (
        '"Google Chrome";v="153", "Not_A Brand";v="8", '
        '"Chromium";v="153"'
    ),
    "sec-ch-ua-mobile": "?0",
    "sec-ch-ua-platform": '"Windows"',
}

DETAIL_HEADERS = {
    key: value
    for key, value in HEADERS.items()
    if key not in {"Content-Type", "Origin"}
}

# 输出字段：主要岗位字段以及后续入库、核验需要的业务标识。
OUTPUT_FIELDS = [
    ("岗位名称", "bce055"),
    ("企业", "aab004"),
    ("行业", "aab022Name"),
    ("地区", "acb204Name"),
    ("技能要求", "aac015Name"),
    ("岗位要求", "_position_requirements"),
    ("薪资范围", "bcca68"),
    ("发布时间", "bdb286"),
    ("职位ID", "bcb009"),
    ("企业ID", "aab001"),
    ("企业详情ID", "bbb911"),
    ("学历要求", "aac011Name"),
    ("工作经验要求", "aae162Name"),
    ("工作性质", "acb239Name"),
    ("招聘人数", "acb240"),
    ("福利待遇", "bcb182Name"),
    ("信息来源", "bze433Name"),
    ("信息来源代码", "bze433"),
    ("是否在招", "_is_active"),
]

COLUMN_WIDTHS = {
    "岗位名称": 28,
    "企业": 38,
    "行业": 34,
    "地区": 30,
    "技能要求": 22,
    "岗位要求": 60,
    "薪资范围": 20,
    "发布时间": 20,
    "职位ID": 24,
    "企业ID": 24,
    "企业详情ID": 24,
    "学历要求": 18,
    "工作经验要求": 18,
    "工作性质": 14,
    "招聘人数": 12,
    "福利待遇": 50,
    "信息来源": 28,
    "信息来源代码": 16,
    "是否在招": 12,
}


def build_request_data(
    page: int,
    province_code: str,
    province_name: str,
) -> dict[str, Any]:
    """构造岗位列表接口的请求参数。"""
    return {
        "bce055": "",
        "acb241": "",
        "acb242": "",
        "aab056": "",
        "lately": -1,
        "bze433": "",
        "aac015": "",
        "aac011": "",
        "aae162": "",
        "acb239": "",
        "acb204": province_code,
        "acb204Name": province_name,
        "gzxz": "0",
        "pageTag": "01",
        "current": page,
        "size": POSITION_PAGE_SIZE,
        "orderType": "01",
    }


async def get_positions(
    page: int,
    province_code: str,
    province_name: str,
) -> tuple[int, list[dict[str, Any]]]:
    """获取一页岗位数据，并返回页码和记录列表。"""
    response = await async_request_with_proxy(
        method="POST",
        url=URL,
        headers=HEADERS,
        json=build_request_data(page, province_code, province_name),
    )

    data = response.json()
    if data.get("appcode") != "200" or data.get("success") is not True:
        raise RuntimeError(
            f"{province_name}第 {page} 页业务状态异常："
            f"appcode={data.get('appcode')}, success={data.get('success')}"
        )

    records = data.get("data", {}).get("records", [])
    if not isinstance(records, list):
        raise RuntimeError(f"{province_name}第 {page} 页 records 不是列表")

    return page, records


def extract_position(record: dict[str, Any]) -> dict[str, Any]:
    """从原始记录中提取需要写入 Excel 的字段。"""
    result = {
        output_name: record.get(source_name)
        for output_name, source_name in OUTPUT_FIELDS
    }

    # 刚从在招岗位列表抓取到的数据，初始化为1；后续入库后由定时任务更新。
    result["是否在招"] = 1
    result["岗位要求"] = ""

    # 招聘人数转为数字，方便在 Excel 中求和。
    try:
        result["招聘人数"] = int(result["招聘人数"] or 0)
    except (TypeError, ValueError):
        result["招聘人数"] = 0

    return result


def make_position_key(record: dict[str, Any]) -> str:
    """优先用职位ID去重；没有职位ID时使用业务字段组成备用键。"""
    position_id = record.get("bcb009")
    if position_id:
        return f"id:{position_id}"

    fallback = (
        record.get("aab001"),
        record.get("aab004"),
        record.get("bce055"),
        record.get("acb204"),
        record.get("bdb286"),
    )
    return "fallback:" + "|".join(str(value or "") for value in fallback)


async def crawl_positions(
    province_code: str,
    province_name: str,
) -> tuple[list[dict[str, Any]], int]:
    """每批并发抓取10页，遇到最后一页后停止。"""
    unique_positions: dict[str, dict[str, Any]] = {}
    raw_count = 0
    start_page = 1

    while start_page <= POSITION_MAX_PAGES:
        page_numbers = list(
            range(
                start_page,
                min(
                    start_page + POSITION_BATCH_SIZE,
                    POSITION_MAX_PAGES + 1,
                ),
            )
        )

        outcomes = await asyncio.gather(
            *(
                get_positions(page, province_code, province_name)
                for page in page_numbers
            ),
            return_exceptions=True,
        )

        page_records: dict[int, list[dict[str, Any]]] = {}
        failed_pages: list[str] = []

        for requested_page, outcome in zip(page_numbers, outcomes):
            if isinstance(outcome, BaseException):
                failed_pages.append(f"第{requested_page}页：{outcome}")
                continue

            page, records = outcome
            page_records[page] = records

        if failed_pages:
            raise RuntimeError(
                "本批次存在抓取失败页面，为避免生成缺页数据已停止：\n"
                + "\n".join(failed_pages)
            )

        terminal_page: int | None = None
        terminal_reason = ""

        # 按页码处理，避免并发完成顺序打乱数据。
        for page in page_numbers:
            records = page_records[page]
            raw_count += len(records)

            for record in records:
                key = make_position_key(record)
                if key not in unique_positions:
                    unique_positions[key] = extract_position(record)

            if terminal_page is None:
                if not records:
                    terminal_page = page
                    terminal_reason = "records 为空"
                elif len(records) < POSITION_PAGE_SIZE:
                    terminal_page = page
                    terminal_reason = (
                        f"仅返回 {len(records)} 条，少于每页 "
                        f"{POSITION_PAGE_SIZE} 条"
                    )

        print(
            f"{province_name}：已完成第 "
            f"{page_numbers[0]}-{page_numbers[-1]} 页，"
            f"原始记录 {raw_count} 条，去重后 {len(unique_positions)} 条"
        )

        if terminal_page is not None:
            later_nonempty_pages = [
                page
                for page in page_numbers
                if page > terminal_page and page_records[page]
            ]
            if later_nonempty_pages:
                raise RuntimeError(
                    f"第 {terminal_page} 页{terminal_reason}，"
                    f"但后续页面仍有数据：{later_nonempty_pages}。"
                    "分页结果可能发生了变化，请重新运行。"
                )

            print(
                f"{province_name}：第 {terminal_page} 页{terminal_reason}，"
                "停止继续翻页。"
            )
            break

        start_page += POSITION_BATCH_SIZE
        if POSITION_BATCH_INTERVAL_SECONDS > 0:
            await asyncio.sleep(POSITION_BATCH_INTERVAL_SECONDS)
    else:
        raise RuntimeError(
            f"已达到安全页数上限 {POSITION_MAX_PAGES}，仍未发现最后一页"
        )

    return list(unique_positions.values()), raw_count


async def get_position_requirements(
    position_id: str,
    semaphore: asyncio.Semaphore,
) -> tuple[str, str | None]:
    """访问公开岗位详情接口，返回acb22a和空值原因。"""
    async with semaphore:
        try:
            response = await async_request_with_proxy(
                method="GET",
                url=POSITION_DETAIL_URL.format(position_id=position_id),
                headers=DETAIL_HEADERS,
            )
        finally:
            if POSITION_DETAIL_REQUEST_INTERVAL_SECONDS > 0:
                await asyncio.sleep(POSITION_DETAIL_REQUEST_INTERVAL_SECONDS)

    payload = response.json()
    if not isinstance(payload, dict):
        raise RuntimeError("岗位详情接口的JSON根节点不是对象")
    if str(payload.get("appcode")) != "200" or payload.get("success") is False:
        raise RuntimeError(
            f"业务状态异常：appcode={payload.get('appcode')}, "
            f"success={payload.get('success')}"
        )

    data = payload.get("data")
    if not isinstance(data, dict):
        raise RuntimeError("返回的data不是对象")

    returned_id = str(data.get("bcb009") or "").strip()
    if not returned_id:
        return "", "岗位详情返回的bcb009为空"
    if returned_id != position_id:
        raise RuntimeError(f"接口返回了其他岗位ID：{returned_id}")

    requirements = data.get("acb22a")
    if requirements is None:
        return "", "岗位详情未提供acb22a"

    requirements_text = str(requirements).strip()
    if not requirements_text:
        return "", "岗位详情中的acb22a为空"
    return requirements_text, None


async def enrich_position_requirements(
    positions: list[dict[str, Any]],
    province_name: str,
) -> None:
    """分批并发获取岗位要求，单个详情失败不会丢弃整省岗位数据。"""
    targets = [position for position in positions if position.get("职位ID")]
    if not targets:
        return

    semaphore = asyncio.Semaphore(POSITION_DETAIL_CONCURRENCY)
    success_count = 0
    empty_count = 0
    failure_count = 0
    failure_samples: list[str] = []

    for start in range(0, len(targets), POSITION_DETAIL_BATCH_SIZE):
        batch = targets[start : start + POSITION_DETAIL_BATCH_SIZE]
        outcomes = await asyncio.gather(
            *(
                get_position_requirements(
                    str(position["职位ID"]),
                    semaphore,
                )
                for position in batch
            ),
            return_exceptions=True,
        )

        for position, outcome in zip(batch, outcomes):
            position_id = str(position["职位ID"])
            if isinstance(outcome, BaseException):
                failure_count += 1
                if len(failure_samples) < 10:
                    failure_samples.append(f"{position_id}：{outcome}")
                continue

            requirements, empty_reason = outcome
            position["岗位要求"] = requirements
            if requirements:
                success_count += 1
            else:
                empty_count += 1
                if empty_reason and len(failure_samples) < 10:
                    failure_samples.append(f"{position_id}：{empty_reason}")

        print(
            f"{province_name}岗位要求：已处理 "
            f"{min(start + len(batch), len(targets))}/{len(targets)}，"
            f"有内容 {success_count}，空值 {empty_count}，失败 {failure_count}"
        )

    if failure_samples:
        print("岗位要求空值/失败示例（最多10条）：")
        for sample in failure_samples:
            print(f"- {sample}")


def clean_excel_value(value: Any) -> Any:
    """移除 Excel 不允许的控制字符。"""
    if value is None:
        return ""
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value).strip()[:32767]
    return value


def save_to_excel(
    rows: list[dict[str, Any]],
    province_name: str,
) -> Path:
    """将岗位结果保存为带筛选和冻结表头的 xlsx 文件。"""
    output_dir = Path(POSITION_OUTPUT_DIR)
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / (
        f"{province_name}招聘岗位_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    )

    workbook = Workbook(write_only=True)
    worksheet = workbook.create_sheet(title="岗位数据")
    worksheet.freeze_panes = "A2"

    headers = [name for name, _ in OUTPUT_FIELDS]
    for index, header in enumerate(headers, 1):
        worksheet.column_dimensions[get_column_letter(index)].width = (
            COLUMN_WIDTHS[header]
        )

    header_cells = []
    for header in headers:
        cell = WriteOnlyCell(worksheet, value=header)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        header_cells.append(cell)
    worksheet.append(header_cells)

    for row in rows:
        cells = []
        for header in headers:
            cell = WriteOnlyCell(
                worksheet,
                value=clean_excel_value(row.get(header)),
            )
            cell.alignment = Alignment(vertical="top")
            cells.append(cell)
        worksheet.append(cells)

    last_column = get_column_letter(len(headers))
    worksheet.auto_filter.ref = f"A1:{last_column}{len(rows) + 1}"
    workbook.save(output_path)
    return output_path.resolve()


async def main() -> None:
    successes: list[tuple[str, int, int, int, Path]] = []
    failures: list[tuple[str, str]] = []

    for index, (province_name, province_code) in enumerate(PROVINCES, 1):
        print("=" * 60)
        print(
            f"[{index}/{len(PROVINCES)}] 开始抓取："
            f"{province_name}（{province_code}）"
        )

        try:
            positions, raw_count = await crawl_positions(
                province_code,
                province_name,
            )
            await enrich_position_requirements(positions, province_name)
            output_path = save_to_excel(positions, province_name)
            vacancy_count = sum(
                int(position.get("招聘人数") or 0)
                for position in positions
            )
            successes.append(
                (
                    province_name,
                    raw_count,
                    len(positions),
                    vacancy_count,
                    output_path,
                )
            )

            print(f"{province_name}原始记录数：{raw_count}")
            print(f"{province_name}去重后职位数：{len(positions)}")
            print(f"{province_name}在招岗位数：{vacancy_count}")
            print(f"Excel 已保存：{output_path}")
        except Exception as exc:
            failures.append((province_name, str(exc)))
            print(f"{province_name}抓取失败，已跳过：{exc}")

    print("=" * 60)
    print(f"全部省份处理完成：成功 {len(successes)} 个，失败 {len(failures)} 个")
    print(f"成功抓取职位总数：{sum(item[2] for item in successes)}")
    print(f"成功抓取在招岗位总数：{sum(item[3] for item in successes)}")

    if failures:
        print("失败省份：")
        for province_name, error in failures:
            print(f"- {province_name}：{error}")


if __name__ == "__main__":
    asyncio.run(main())
